#!/usr/bin/env python
"""
build_smishing_workbook.py -- one Excel workbook that joins the team's links, the team's own labels and
what the isolated lab actually saw (WEB_INSPECTION_PLAN.md sections 12.5, 15). Run on the HOST, after the
batch was pulled and validated:

    set PYTHONPATH=D:\\LeJetData\\devlibs
    python scripts/build_smishing_workbook.py --name smish_oct --batch smish01 ^
        --out D:\\LeJetData\\reports\\smishing_links_cleaned.xlsx

Inputs (all under D:\\LeJetData): input\\<name>.meta.json (from clean_smishing_links.py),
clean\\<batch>.clean.jsonl (validated lab data), clean\\shots\\<batch>\\shots.jsonl (screenshot attempts) and,
if present, clean\\<batch>.model_input.csv (page rank).

Sheets:
  Summary            counts, definitions, and how the team's labels compare with what the lab saw
  Cleaned            every unique link with its sources, the team's label and the lab's result
  Live               links whose page the lab loaded (any team label, "unsure" included)
  Reachable-blocked  the site answered but not with a usable page (403, bot check, parked, ...)
  Not live           no answer: DNS failure, refused or silent server, redirect loop, ...
  Unsure             every link the team labelled "unsure", with the lab's result (nothing is dropped)

EVERY URL in the workbook is DEFANGED (hxxps://evil[.]com). The refanged list stays in D:\\LeJetData\\input.
Reviewer names from the original sheet are not copied.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "src"))

from urllib.parse import urlsplit                              # noqa: E402

from lejet import domains, safety, urlnorm                     # noqa: E402

# plain-language meaning of the lab's fail_reason
WHY = {
    "none": "page loaded",
    "http_4xx": "site answered with an error or 'forbidden' (HTTP 4xx)",
    "http_5xx": "site answered with a server error (HTTP 5xx)",
    "challenge_page": "bot check / challenge page",
    "parked_or_default_page": "parked domain or default server page",
    "blocked_or_empty": "page empty or blocked",
    "non_html": "not a web page (a file or data)",
    "too_large": "page too large to read",
    "tls_error": "certificate error",
    "dns_nxdomain": "domain does not exist",
    "dns_no_answer": "domain has no address",
    "dns_timeout": "name lookup timed out",
    "dns_servfail": "name lookup failed",
    "connect_refused": "server refused the connection",
    "connect_timeout": "server did not answer",
    "read_timeout": "server too slow",
    "too_many_redirects": "redirect loop / too many redirects",
    "redirect_to_blocked": "redirected to a place the lab refuses to visit",
    "rejected_url": "refused by the lab's safety rules",
    "port_blocked": "refused by the lab's safety rules (port)",
    "internal_error": "lab error",
}
ANSWERED = {"http_4xx", "http_5xx", "challenge_page", "parked_or_default_page", "blocked_or_empty", "non_html", "too_large"}


def lab_class(fail_reason: str, screenshot_ok: bool) -> str:
    """live = the page loaded; reachable-blocked = something answered; not live = nothing usable."""
    if fail_reason == "none":
        return "live"
    if fail_reason in ANSWERED or screenshot_ok:
        return "reachable-blocked"
    return "not live"


def refine(m: Dict[str, Any], row: Dict[str, Any], cls: str, why: str) -> Tuple[str, str]:
    """A 404/410 from the SHORTENER itself, or a Cloudflare 52x, is not "blocked": the link was removed (or its
    code is wrong) / the site's server is down. Returns the corrected (class, plain text)."""
    status = row.get("final_status")
    final = str(row.get("final_reg_domain", "")).replace("[.]", ".").lower()
    own = urlnorm.registrable_domain(m["key"].split("/")[0].split("?")[0].split(":")[0]) or ""
    if m.get("shortener") and final and final == own and status in (404, 410):
        return "not live", ("the shortener says the link was removed or its code is wrong (HTTP %s); short codes are "
                            "case-sensitive: compare with the original message" % status)
    if isinstance(status, int) and 520 <= status <= 530:
        return "not live", f"the site's server is down behind its protection service (HTTP {status})"
    return cls, why


def caveats(m: Dict[str, Any], row: Dict[str, Any], rank: str) -> str:
    """Reasons a page that LOADED may still not be a phishing page (or may not describe one). Plain text."""
    out: List[str] = []
    final = str(row.get("final_reg_domain", "")).replace("[.]", ".").lower()
    own = urlnorm.registrable_domain(m["key"].split("/")[0].split("?")[0].split(":")[0]) or ""
    if row.get("fail_reason") == "none":
        if m.get("shortener") and final and final == own:
            out.append("the shortener's own page: the link was probably disabled or removed")
        platform = row.get("hosting_platform") if row.get("free_hosting_subdomain") else ""
        if not row.get("free_hosting_subdomain"):                  # also catch runs made before the list grew
            try:
                final_url = str(row.get("final_url", "")).replace("[.]", ".")
                if final_url.lower().startswith("hxxp"):
                    final_url = "http" + final_url[4:]
                parts = urlsplit(final_url)
                hosting = domains.classify_hosting(domains.clean_host(parts.hostname or ""), parts.path)
                if hosting.free_hosting_subdomain:
                    platform = hosting.hosting_platform or "unknown"
                    row = dict(row, free_hosting_subdomain=1)
            except Exception:                                      # noqa: BLE001 - a caveat is a hint, never a crash
                pass
        if row.get("free_hosting_subdomain"):
            out.append(f"free hosting platform ({platform or 'unknown'}): "
                       "domain age and rank belong to the platform")
        try:
            if rank != "" and int(rank) >= 9:
                out.append("very popular destination (page rank 9-10): probably not a phishing page")
        except ValueError:
            pass
    return "; ".join(out)


def team_label(manual: List[str]) -> str:
    if not manual:
        return "(no label)"
    return "/".join(manual)


def read_jsonl(path: str) -> List[Dict[str, Any]]:
    with open(path, encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]


def join(meta: List[Dict[str, Any]], rows: List[Dict[str, Any]], shots: Dict[str, Dict[str, Any]],
         ranks: Dict[str, str]) -> List[Dict[str, Any]]:
    """One record per unique link: the team's side plus the lab's side."""
    by_id = {r["url_id"]: r for r in rows}
    out: List[Dict[str, Any]] = []
    for i, m in enumerate(meta, 1):
        uid = urlnorm.compute_url_id(m["raw"])
        row = by_id.get(uid)
        shot = shots.get(uid)
        rec: Dict[str, Any] = {
            "n": i, "url": safety.defang_url(m["raw"]), "shortener": "yes" if m.get("shortener") else "",
            "sources": "; ".join(sorted({s.split(":")[0] + " row " + s.rsplit(":", 1)[-1] for s in m["sources"]})),
            "team": team_label(m.get("manual", [])), "url_id": uid, "note": m.get("note", ""),
        }
        if row is None:
            rec.update(lab="not tested", why=m.get("note") or "not in the lab batch", fail="", final_url="", final_domain="",
                       status="", hops="", shot="", page_rank="", domain_age="", cls="not live", caveat="")
        else:
            shot_ok = bool(shot and shot.get("ok"))
            fail = row["fail_reason"]
            age = row.get("features", {}).get("domain_age", -1)
            cls, why = refine(m, row, lab_class(fail, shot_ok), WHY.get(fail, fail))
            rec.update(lab=WHY.get(fail, fail), why=why, fail=fail,
                       final_url=row.get("final_url", ""), final_domain=row.get("final_reg_domain", ""),
                       status=row.get("final_status", ""), hops=row.get("hop_count", ""),
                       shot=("yes" if shot_ok else ("no: " + (shot or {}).get("error", "not attempted"))),
                       page_rank=ranks.get(uid, ""), domain_age=("" if age in (-1, None) else age),
                       cls=cls, caveat=caveats(m, row, ranks.get(uid, "")))
            if shot_ok and shot:                                    # where the BROWSER ended up, and how much it shows
                # page-controlled text: any link inside it is neutralised again, whatever the VM sent
                rec["browser_url"] = safety.defang_embedded_urls(shot.get("final_url", ""))
                rec["title"] = safety.defang_embedded_urls(shot.get("title", ""))
                if shot.get("rescued"):
                    rec["blank"] = "no: picture taken before the page wiped itself (hides from automated browsers)"
                elif shot.get("text_chars") == 0 and shot.get("bytes", 0) < 12000:
                    rec["blank"] = "yes: white screen (hides from automated browsers)"
                else:
                    rec["blank"] = ""
        out.append(rec)
    return out


COLUMNS = [("n", "#", 5), ("url", "Link (defanged)", 52), ("shortener", "Shortener", 10), ("team", "Team's label", 13),
           ("cls", "Lab result", 18), ("why", "What the lab saw", 44), ("caveat", "Caveat (check before using)", 52),
           ("final_url", "Final address (defanged)", 52),
           ("final_domain", "Final domain", 28), ("status", "HTTP status", 11), ("hops", "Redirects", 10),
           ("page_rank", "Page rank (0-10)", 15), ("domain_age", "Domain age (days)", 17), ("shot", "Screenshot", 22),
           ("blank", "White screen?", 18), ("browser_url", "Where the browser ended (defanged)", 52),
           ("title", "Page title (as shown)", 40),
           ("url_id", "Screenshot file / id", 20), ("sources", "Found in", 34)]


def agreement(recs: List[Dict[str, Any]]) -> List[List[Any]]:
    labels = sorted({r["team"] for r in recs})
    classes = ["live", "reachable-blocked", "not live"]
    table = [["Team's label \\ Lab result"] + classes + ["total"]]
    for lab in labels:
        sub = [r for r in recs if r["team"] == lab]
        table.append([lab] + [sum(1 for r in sub if r["cls"] == c) for c in classes] + [len(sub)])
    table.append(["total"] + [sum(1 for r in recs if r["cls"] == c) for c in classes] + [len(recs)])
    return table


def write_workbook(recs: List[Dict[str, Any]], out: str, info: Dict[str, str]) -> None:
    import openpyxl                                              # lazy: host helper, not a repo dependency
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(title: str, subset: List[Dict[str, Any]]) -> None:
        ws = wb.create_sheet(title)
        ws.append([c[1] for c in COLUMNS])
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r in subset:
            ws.append([r.get(k, "") for k, _, _ in COLUMNS])
        for i, (_, _, width) in enumerate(COLUMNS, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        # a text cell never becomes a formula or a link
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

    ws = wb.active
    ws.title = "Summary"
    cls = Counter(r["cls"] for r in recs)
    ws.append(["Smishing links: cleaned and checked in the isolated lab"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in info.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Unique links", len(recs)])
    ws.append(["  Live (the page loaded)", cls["live"]])
    ws.append(["  Reachable but blocked (answered, no usable page)", cls["reachable-blocked"]])
    ws.append(["  Not live (no answer)", cls["not live"]])
    ws.append(["Screenshots taken", sum(1 for r in recs if r["shot"] == "yes")])
    ws.append(["Link shorteners among the links", sum(1 for r in recs if r["shortener"])])
    ws.append([])
    ws.append(["How the team's labels compare with what the lab saw (counts of links)"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for line in agreement(recs):
        ws.append(line)
    ws.append([])
    for note in (
        "Live = the lab fetched the page and could compute the features. Reachable-blocked = a server answered but not with a usable "
        "page (403, bot check, parked page ...) or the browser got a page the plain fetch could not. Not live = DNS failure, no "
        "answer, a redirect loop, a server that is down, or a shortener that says the link was removed (HTTP 404/410: its codes are "
        "case-sensitive, so compare them with the original message).",
        "'White screen?' marks pages whose picture is blank and that show no text at all: they were reachable, but a script-built "
        "or phone-only page may need the phone profile (Android browser) to show anything.",
        "The team's labels are a human judgement from the day the links were reviewed; the lab result is from the run date above. "
        "Smishing sites are short-lived, so a link can honestly change from live to dead in between.",
        "Several sites refuse the VPN provider's address range (the screenshot shows the block page), so 'reachable-blocked' does not "
        "mean the site is harmless or gone.",
        "Every link in this workbook is defanged (hxxps://example[.]com). Do not open links from it in a browser.",
    ):
        ws.append([note])
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 18
    for col in "CDE":
        ws.column_dimensions[col].width = 20

    sheet("Cleaned", recs)
    sheet("Live", [r for r in recs if r["cls"] == "live"])
    sheet("Reachable-blocked", [r for r in recs if r["cls"] == "reachable-blocked"])
    sheet("Not live", [r for r in recs if r["cls"] == "not live"])
    sheet("Unsure", [r for r in recs if "unsure" in r["team"]])
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--name", required=True, help="the cleaned list name (input\\<name>.meta.json)")
    ap.add_argument("--batch", required=True)
    ap.add_argument("--data", default=r"D:\LeJetData")
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)

    meta = json.load(open(os.path.join(a.data, "input", f"{a.name}.meta.json"), encoding="utf-8"))
    clean = os.path.join(a.data, "clean")
    rows = read_jsonl(os.path.join(clean, f"{a.batch}.clean.jsonl"))
    shots: Dict[str, Dict[str, Any]] = {}
    shots_path = os.path.join(clean, "shots", a.batch, "shots.jsonl")
    if os.path.exists(shots_path):
        shots = {r["url_id"]: r for r in read_jsonl(shots_path)}
    ranks: Dict[str, str] = {}
    mi = os.path.join(clean, f"{a.batch}.model_input.csv")
    if os.path.exists(mi):
        with open(mi, encoding="utf-8", newline="") as fh:
            for r in csv.DictReader(fh):
                if r.get("page_rank_ok") == "1":
                    ranks[r["url_id"]] = r["page_rank"]
    recs = join(meta, rows, shots, ranks)
    info = {"Batch": a.batch, "Lab run": rows[0]["fetch_ts"][:10] if rows else "", "Vantage point": "Mullvad VPN exit, Manila (PH)",
            "Source files": "asdasdasdas.xlsx and linkss.xlsx (merged, duplicates removed)",
            "Screenshots": f"D:\\LeJetData\\clean\\shots\\{a.batch}\\index.html"}
    write_workbook(recs, a.out, info)
    cls = Counter(r["cls"] for r in recs)
    print(f"{len(recs)} unique links: live {cls['live']}, reachable-blocked {cls['reachable-blocked']}, not live {cls['not live']}")
    print(f"wrote {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
