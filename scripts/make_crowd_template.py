#!/usr/bin/env python
"""
make_crowd_template.py -- write the Excel template the crowdsourced SMS are tagged in, with a
dropdown on every column that has fixed answers (docs/c_prune_architecture.md section 4.2).
Run on the HOST:

    python scripts/make_crowd_template.py
    python scripts/make_crowd_template.py --out data/crowd/raw/_template_submissions.xlsx

Copy the template to submissions_<YYYY-MM-DD>_<code>.xlsx (one per person per day); crowd_to_urls.py
reads that file directly. The dropdown lists sit on a hidden "lists" sheet. Excel only warns on
a value that is not in the list (it does not block it), so a pasted row with a typo is still kept
and can be fixed by hand.

submitter_view, org_type and concern use the Google Form's own wording, so answers can be
copied straight from the form export.
"""

from __future__ import annotations

import argparse
import datetime
import os
import sys
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from crowd_to_urls import COLUMNS, EXAMPLE_TEXT, SHEET   # noqa: E402

DEFAULT_OUT = os.path.join(HERE, "..", "data", "crowd", "raw", "_template_submissions.xlsx")
ROWS = 2000   # rows that get the dropdowns and date formats

CHOICES: Dict[str, List[str]] = {
    "submitter_view": ["Suspicious or potentially fraudulent", "Appears legitimate or genuine", "Not sure"],
    "org_type": ["Bank", "E-wallet / digital payment service", "Government agency",
                 "Telecommunications provider", "Delivery / courier service",
                 "Online shopping / e-commerce platform", "Loan / lending service",
                 "Credit card / financial service", "Utility / billing service", "Unknown sender", "Other"],
    "concern": ["Account verification", "Account suspension / blocking", "Login or password",
                "OTP / authentication", "Payment or transaction", "Fine / penalty / violation", "Toll payment",
                "Loan / debt", "Refund", "Prize / reward", "Delivery / parcel fee", "Credit / debit card",
                "Investment", "Tax / government payment", "Promotion / advertisement", "Other financial matter",
                "None of the above", "Not sure"],
    "label": ["smishing", "legitimate"],
    "category": ["bank", "e-wallet", "credit_card", "loan", "bill_payment", "delivery_fee", "government_aid",
                 "investment", "other_financial"],
    "language": ["en", "tl", "taglish"],
    "sender_type": ["number", "shortcode", "alphanumeric", "unknown"],
    "label_source": ["submitter", "bank_confirmed", "reviewer"],
    "consent": ["1", "0"],
}

# Shown when a header cell is hovered (and as the dropdown's input message).
HELP: Dict[str, str] = {
    "date_submitted": "When the form was answered, Philippine time: YYYY-MM-DD HH:MM",
    "sms_text": "The full SMS copied exactly, link included. Names/numbers may stay here; the builder masks them.",
    "submitter_view": "Form question: How would you describe the SMS message?",
    "org_type": "Form question: What type of sender or organization does the message appear to come from?",
    "concern": "Form question: What does the message mainly concern?",
    "label": "The checked decision. Never change it because the page was dead or looked clean.",
    "category": "Financial category used by the study dataset.",
    "language": "en = English, tl = Tagalog, taglish = mixed",
    "sender_type": "number = 11-digit mobile, shortcode = e.g. 22565, alphanumeric = a name like GCash. Do NOT write the sender itself.",
    "date_received": "The day the person got the SMS: YYYY-MM-DD",
    "label_source": "Who gave the label.",
    "label_verified_by": "Code of the SECOND person who checked the label, e.g. r02. Empty until checked.",
    "consent": "1 = agreed, 0 = did not (0 rows are dropped).",
}

WIDTHS = {"date_submitted": 17, "sms_text": 60, "submitter_view": 30, "org_type": 32, "concern": 28,
          "label": 12, "category": 16, "language": 10, "sender_type": 13, "date_received": 13,
          "label_source": 15, "label_verified_by": 16, "consent": 9}

EXAMPLE = {"date_submitted": datetime.datetime(2026, 11, 3, 14, 20), "sms_text": EXAMPLE_TEXT,
           "submitter_view": "Suspicious or potentially fraudulent", "org_type": "E-wallet / digital payment service",
           "concern": "Account suspension / blocking", "label": "smishing", "category": "e-wallet",
           "language": "en", "sender_type": "number", "date_received": datetime.date(2026, 11, 3),
           "label_source": "submitter", "label_verified_by": "r02", "consent": "1"}


def build() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    lists = wb.create_sheet("lists")
    lists.sheet_state = "hidden"

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill = head_font, head_fill
        cell.comment = Comment(HELP[name], "Le-Jet")
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[name]
    ws.freeze_panes = "C2"   # header and the first two columns stay visible while scrolling
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    list_col = 0
    for c, name in enumerate(COLUMNS, start=1):
        letter = get_column_letter(c)
        rng = f"{letter}2:{letter}{ROWS + 1}"
        if name in CHOICES:
            # The list lives on the hidden sheet under a name, so long lists (concern has 18) still work.
            list_col += 1
            lc = get_column_letter(list_col)
            values = CHOICES[name]
            lists.cell(row=1, column=list_col, value=name)
            for r, v in enumerate(values, start=2):
                lists.cell(row=r, column=list_col, value=v)
            wb.defined_names[f"list_{name}"] = DefinedName(
                f"list_{name}", attr_text=f"lists!${lc}$2:${lc}${len(values) + 1}")
            dv = DataValidation(type="list", formula1=f"=list_{name}", allow_blank=True, showErrorMessage=True,
                                errorStyle="warning", errorTitle="Not in the list",
                                error="This value is not one of the choices. Keep it anyway?",
                                promptTitle=name, prompt=HELP[name][:255], showInputMessage=True)
            dv.add(rng)
            ws.add_data_validation(dv)
        if name == "date_submitted":
            fmt = "yyyy-mm-dd hh:mm"
        elif name == "date_received":
            fmt = "yyyy-mm-dd"
        else:
            fmt = "@"   # text, so codes like r02 and 1 are kept as typed
        for r in range(2, ROWS + 2):
            ws.cell(row=r, column=c).number_format = fmt
            if name == "sms_text":
                ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    for c, name in enumerate(COLUMNS, start=1):
        ws.cell(row=2, column=c, value=EXAMPLE[name])
    return wb


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", default=DEFAULT_OUT)
    a = ap.parse_args(argv)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    build().save(a.out)
    print(f"wrote {a.out} ({len(COLUMNS)} columns, dropdowns on {len(CHOICES)}, {ROWS} rows ready)")
    print("row 2 is the example (fake GCash message): delete it in your copy")
    return 0


if __name__ == "__main__":
    sys.exit(main())
